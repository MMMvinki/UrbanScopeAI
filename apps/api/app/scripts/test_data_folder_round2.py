"""数据目录接入 · 第二轮回归测试（内里修复后新增验证）。

覆盖（对应修复项）：
12. GBK/GB18030 CSV：扫描识别 + 启用成功（曾因 _read_csv 硬编码 utf-8-sig 抛 UnicodeDecodeError）；
13. 矢量要素数门槛：文件小但要素极多 → 仅可查询；预览启用走 fiona 流式抽稀（不全量读内存）；
14. 大型 XLSX：按字节判仅可查询 → 预览启用（曾走 gpd.read_file 读不了 xlsx 直接崩）；
15. 跨数据源 dataset_id 去重：两个目录同名文件启用不互相覆盖。

运行：cd apps/api && .venv/Scripts/python.exe app/scripts/test_data_folder_round2.py
"""

import json
import os
import shutil
import sys
import tempfile
from pathlib import Path

# 必须在 import app.* 之前设置 env（模块级常量读取）。
_WS = tempfile.mkdtemp(prefix="data_folder_r2_ws_")
_FIX = tempfile.mkdtemp(prefix="data_folder_r2_fix_")
os.environ["URBANSCOPE_WORKSPACE"] = _WS
# 字节阈值放宽：本轮聚焦「要素数门槛 / 编码 / 流式」，不按字节模拟大文件
os.environ["URBANSCOPE_DATAFOLDER_ROWCOUNT_MAX_BYTES"] = "100000000"
os.environ["URBANSCOPE_DATAFOLDER_ENABLE_MAX_BYTES"] = "100000000"
os.environ["URBANSCOPE_DATAFOLDER_ENABLE_MAX_FEATURES"] = "100000"
# 测试用小值：>2000 要素的矢量 → 仅可查询
os.environ["URBANSCOPE_DATAFOLDER_ENABLE_MAX_FEATURES_VECTOR"] = "2000"
os.environ["URBANSCOPE_DATAFOLDER_PREVIEW_MAX_FEATURES"] = "50000"

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

import app.data_folder as df  # noqa: E402
from app.data_folder import DataFolderError  # noqa: E402
from app.data_registry import DATA_REGISTRY  # noqa: E402
from app.models.data_folder import DatasetStatus  # noqa: E402

FAILS: list[str] = []
TOTAL = 0


def check(name: str, cond: bool, detail: str = "") -> None:
    global TOTAL
    TOTAL += 1
    mark = "✓" if cond else "✗"
    print(f"  {mark} {name}" + (f"  —— {detail}" if detail else ""))
    if not cond:
        FAILS.append(name)


def make_shp(dir_: Path, stem: str, n: int = 2, epsg: int = 4326) -> Path:
    import fiona
    from fiona.crs import CRS

    shp = dir_ / f"{stem}.shp"
    schema = {"geometry": "Point", "properties": {"v": "int", "name": "str"}}
    with fiona.open(str(shp), "w", driver="ESRI Shapefile", crs=CRS.from_epsg(epsg), schema=schema, encoding="utf-8") as dst:
        for i in range(n):
            dst.write({"geometry": {"type": "Point", "coordinates": (114.0 + i * 0.001, 22.5)}, "properties": {"v": i, "name": f"p{i}"}})
    return shp


def main() -> int:
    print("== 数据目录接入 · 第二轮回归 ==")

    # ── 12. GBK CSV ────────────────────────────────────────────
    gbk_dir = Path(_FIX) / "gbk"
    gbk_dir.mkdir(parents=True, exist_ok=True)
    gbk_csv = gbk_dir / "gbk.csv"
    gbk_csv.write_bytes("lon,lat,名称,类型\n114.0,22.5,深圳市,交通\n114.1,22.6,广州市,住宿\n".encode("gb18030"))
    gbk_src = df.link_source(str(gbk_dir), alias="gbk源")
    gbk_ds = next((d for d in gbk_src.datasets if d.dataset_id == "gbk"), None)
    check("12a GBK CSV 扫描识别为可启用", gbk_ds is not None and gbk_ds.status == DatasetStatus.enableable)
    r = df.enable_dataset(gbk_src.source_id, "gbk")
    lyr = DATA_REGISTRY.get_layer("gbk")
    ok = r.get("layer_id") == "gbk" and lyr is not None and lyr.feature_count == 2
    # 属性应正确解码（GBK 中文不被破坏成乱码）
    fc = DATA_REGISTRY._geojson.get("gbk", {}).get("features", [])
    names = {f.get("properties", {}).get("名称") for f in fc if "名称" in (f.get("properties") or {})}
    ok = ok and names == {"深圳市", "广州市"}
    check("12b GBK CSV 启用成功且中文属性解码", ok, f"names={names}")

    # ── 13. 矢量要素数门槛 + fiona 流式抽稀 ───────────────────
    vec_dir = Path(_FIX) / "vec"
    vec_dir.mkdir(parents=True, exist_ok=True)
    make_shp(vec_dir, "many", n=2500)   # 2500 > 2000 门槛 → 仅可查询
    make_shp(vec_dir, "few", n=10)      # 10 < 2000 → 可直接启用
    vec_ds = df.scan_folder(vec_dir)
    many = next((d for d in vec_ds if d.dataset_id == "many"), None)
    few = next((d for d in vec_ds if d.dataset_id == "few"), None)
    check("13a 要素数超门槛 → 仅可查询", many is not None and many.status == DatasetStatus.query_only)
    check("13b 要素数低于门槛 → 可直接启用", few is not None and few.status == DatasetStatus.enableable)
    vec_src = df.link_source(str(vec_dir), alias="vec源")
    r_many = df.enable_dataset(vec_src.source_id, "many")  # 走 _ingest_preview → _sample_vector_to_fc
    check("13c 大 shp 预览抽稀要素数受控", r_many.get("feature_count", 0) <= 50000 and r_many.get("feature_count", 0) > 0,
          f"fc={r_many.get('feature_count')}")

    # ── 14. 大型 XLSX 预览启用 ────────────────────────────────
    df._ENABLE_MAX_BYTES = 1000  # 临时收紧字节阈值，让 3000 行 xlsx 变「仅可查询」（读模块全局，运行时生效）
    xl_dir = Path(_FIX) / "xl"
    xl_dir.mkdir(parents=True, exist_ok=True)
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["lon", "lat", "name"])
    for i in range(3000):
        ws.append([114.0 + i * 0.001, 22.5, f"p{i}"])
    xl_path = xl_dir / "big.xlsx"
    wb.save(xl_path)
    xl_src = df.link_source(str(xl_dir), alias="xl源")
    xl_ds = next((d for d in xl_src.datasets if d.dataset_id == "big"), None)
    check("14a 大 XLSX 标仅可查询", xl_ds is not None and xl_ds.status == DatasetStatus.query_only)
    r_xl = df.enable_dataset(xl_src.source_id, "big")  # 走 _ingest_preview → _sample_xlsx_to_fc
    check("14b 大 XLSX 预览启用成功", r_xl.get("feature_count", 0) > 0 and r_xl.get("feature_count", 0) <= 50000,
          f"fc={r_xl.get('feature_count')}")
    df._ENABLE_MAX_BYTES = 100000000  # 恢复

    # ── 15. 跨数据源 dataset_id 去重 ──────────────────────────
    s1 = Path(_FIX) / "s1"
    s2 = Path(_FIX) / "s2"
    s1.mkdir(parents=True, exist_ok=True)
    s2.mkdir(parents=True, exist_ok=True)
    make_shp(s1, "buildings", n=5)
    make_shp(s2, "buildings", n=5)
    src1 = df.link_source(str(s1), alias="源一")
    src2 = df.link_source(str(s2), alias="源二")
    b1 = next((d for d in src1.datasets if d.dataset_id == "buildings"), None)
    b2 = next((d for d in src2.datasets if d.dataset_id.startswith("buildings")), None)
    check("15a 第二目录同名数据集加后缀去重", b1 is not None and b2 is not None and b2.dataset_id != "buildings",
          f"b2.id={b2.dataset_id if b2 else None}")
    r1 = df.enable_dataset(src1.source_id, "buildings")
    r2 = df.enable_dataset(src2.source_id, b2.dataset_id)
    check("15b 两源同名启用为不同图层", r1.get("layer_id") != r2.get("layer_id"),
          f"{r1.get('layer_id')} vs {r2.get('layer_id')}")

    # ── 16. 行政区级 OD/来源表识别（如百度慧眼「客流来源」）────────
    od_dir = Path(_FIX) / "od"
    od_dir.mkdir(parents=True, exist_ok=True)
    od_txt = od_dir / "客流来源分小时.txt"
    od_txt.write_text("区域名\t日期\t小时\t来源省份\t来源城市\t来源区县\t人数\n研究范围\t20250328\t21\t广东省\t深圳市\t宝安区\t6653\n", encoding="utf-8")
    od_ds = df.scan_folder(od_dir)
    od = next((d for d in od_ds if d.dataset_id.startswith("客流来源")), None)
    check("16a OD/来源表走表格图层且 warning 标注",
          od is not None and od.status in (DatasetStatus.enableable, DatasetStatus.query_only)
          and od.geometry == "table" and "OD" in " ".join(od.warnings),
          f"status={od.status if od else None} geometry={od.geometry if od else None} warnings={od.warnings if od else None}")

    # ── 17. 无坐标 CSV → 表格图层（可启用、属性保值、可 query）────
    from app.models.layer import GeometryType  # noqa: E402

    tab_dir = Path(_FIX) / "tab"
    tab_dir.mkdir(parents=True, exist_ok=True)
    (tab_dir / "tab17.csv").write_text(
        "name,zip,big_id\n深圳市,0755,12345678901234567890\n广州市,020,99999999999999999999\n",
        encoding="utf-8",
    )
    tab_src = df.link_source(str(tab_dir), alias="表格源")
    tab_ds = next((d for d in tab_src.datasets if d.dataset_id == "tab17"), None)
    check("17a 无坐标 CSV 标表格+可启用",
          tab_ds is not None and tab_ds.status == DatasetStatus.enableable and tab_ds.geometry == "table")
    r17 = df.enable_dataset(tab_src.source_id, "tab17")
    lyr17 = DATA_REGISTRY.get_layer("tab17")
    check("17b 启用为 Table 图层（crs N/A + lazy）",
          lyr17 is not None and lyr17.geometry == GeometryType.table and lyr17.crs == "N/A" and lyr17.lazy)
    check("17c fields 非空 + feature_count=2",
          lyr17 is not None and len(lyr17.fields) > 0 and lyr17.feature_count == 2)
    fc17 = DATA_REGISTRY._geojson.get("tab17", {}).get("features", [])
    props17 = [f.get("properties", {}) for f in fc17]
    zip_ok = any(p.get("zip") == "0755" for p in props17)
    big_ok = any(p.get("big_id") == "12345678901234567890" for p in props17)
    check("17d 前导零/超大 ID 保值", zip_ok and big_ok, f"props={props17}")
    from app.tools.query_features import run_query_features  # noqa: E402

    q = run_query_features({"layer_id": "tab17", "filter": {"name": "深圳市"}})
    check("17e query_features 带 filter 命中", q.success and q.output.get("total_matched") == 1,
          f"msg={q.message}")

    # ── 18. 大无坐标 CSV → 仅可查询 → 抽稀预览 ─────────────────
    df._ENABLE_MAX_BYTES = 1000  # 临时收紧，让 3000 行无坐标 CSV 变 query_only
    bigtab_dir = Path(_FIX) / "bigtab"
    bigtab_dir.mkdir(parents=True, exist_ok=True)
    bigtab_rows = ["name,val"] + [f"p{i},{i}" for i in range(3000)]
    (bigtab_dir / "bigtab.csv").write_text("\n".join(bigtab_rows), encoding="utf-8")
    bigtab_src = df.link_source(str(bigtab_dir), alias="大表格源")
    bigtab_ds = next((d for d in bigtab_src.datasets if d.dataset_id == "bigtab"), None)
    check("18a 大无坐标 CSV 标仅可查询",
          bigtab_ds is not None and bigtab_ds.status == DatasetStatus.query_only and bigtab_ds.geometry == "table")
    r18 = df.enable_dataset(bigtab_src.source_id, "bigtab")
    lyr18 = DATA_REGISTRY.get_layer("bigtab")
    check("18b 大表格抽稀预览（fc≤上限 + 图层名含抽稀预览）",
          r18.get("feature_count", 0) <= 50000 and r18.get("feature_count", 0) > 0
          and "抽稀预览" in (lyr18.name if lyr18 else ""),
          f"fc={r18.get('feature_count')} name={lyr18.name if lyr18 else None}")
    df._ENABLE_MAX_BYTES = 100000000  # 恢复

    # ── 20. disable_dataset 全流程 ─────────────────────────────
    from app.workspace import LAYER_DATA_DIR  # noqa: E402

    r_d = df.disable_dataset(tab_src.source_id, "tab17")
    check("20a 停用后图层移除", DATA_REGISTRY.get_layer("tab17") is None)
    ds17 = next(d for d in df.get_source(tab_src.source_id).datasets if d.dataset_id == "tab17")
    check("20b 停用后 enabled False + layer_id None", not ds17.enabled and ds17.layer_id is None)
    check("20c 磁盘 JSON 删除", not (LAYER_DATA_DIR / "tab17.json").exists())
    r_d2 = df.disable_dataset(tab_src.source_id, "tab17")
    check("20d 第二次停用幂等成功", r_d2.get("already_disabled") is True)
    r_again = df.enable_dataset(tab_src.source_id, "tab17")
    check("20e 停用后可再启用", r_again.get("layer_id") == "tab17" and DATA_REGISTRY.get_layer("tab17") is not None)

    # ── 21. 文件夹级全部启用（硬上限 20 + clamp）────────────
    many_dir = Path(_FIX) / "many"
    many_dir.mkdir(parents=True, exist_ok=True)
    for i in range(25):
        (many_dir / f"d{i}.csv").write_text("name,lon,lat\np,114.0,22.5\n", encoding="utf-8")
    many_src = df.link_source(str(many_dir), alias="批量源")
    r_all = df.enable_all_datasets(many_src.source_id)
    check("21a 硬上限 20", r_all.get("planned_count") == 20 and r_all.get("enabled") == 20
          and r_all.get("capped") is True and r_all.get("capped_by_limit") is True,
          f"planned={r_all.get('planned_count')} enabled={r_all.get('enabled')}")
    df.disable_all_datasets(many_src.source_id)
    r_clamp = df.enable_all_datasets(many_src.source_id, limit=100)
    check("21b limit>20 被 clamp 到 20", r_clamp.get("planned_count") == 20,
          f"planned={r_clamp.get('planned_count')}")

    # ── 22. 文件夹级全部停用 ─────────────────────────────────
    r_disall = df.disable_all_datasets(many_src.source_id)
    check("22a disable-all 停用 20 个", r_disall.get("count") == 20, f"count={r_disall.get('count')}")
    check("22b 全 enabled False", all(not d.enabled for d in df.get_source(many_src.source_id).datasets))

    # ── 23. 栅格元数据 + 缩略图 + chunked 分支 ────────────────
    import numpy as np  # noqa: E402
    import rasterio  # noqa: E402
    from rasterio.transform import from_origin  # noqa: E402

    tif_dir = Path(_FIX) / "tif"
    tif_dir.mkdir(parents=True, exist_ok=True)
    tif_path = tif_dir / "small.tif"
    with rasterio.open(tif_path, "w", driver="GTiff", height=64, width=64, count=3,
                       dtype="uint8", crs="EPSG:4326", transform=from_origin(114.0, 22.6, 0.001, 0.001)) as dst:
        dst.write(np.zeros((3, 64, 64), dtype="uint8"))
    tif_ds = df._probe_tif(tif_path, "small.tif", 0)
    check("23a tif 元数据", tif_ds.raster is not None and tif_ds.raster.width == 64
          and tif_ds.raster.height == 64 and tif_ds.raster.bands == 3,
          f"raster={tif_ds.raster is not None}")
    check("23b tif 缩略图 data-url", tif_ds.raster is not None
          and (tif_ds.raster.preview_image or "").startswith("data:image/png"))
    df._CHUNKED_MIN_PIXELS = 0  # 强制走 chunked 分支
    tif_ds2 = df._probe_tif(tif_path, "small.tif", 0)
    check("23c chunked 分支仍生成缩略图", tif_ds2.raster is not None
          and (tif_ds2.raster.preview_image or "").startswith("data:image/png"))
    df._CHUNKED_MIN_PIXELS = 20_000_000  # 恢复
    with rasterio.open(tif_dir / "gray.tif", "w", driver="GTiff", height=32, width=32, count=1,
                       dtype="uint16", crs=None, transform=from_origin(0, 0, 1, 1)) as dst:
        dst.write(np.zeros((1, 32, 32), dtype="uint16"))
    gray_ds = df._probe_tif(tif_dir / "gray.tif", "gray.tif", 0)
    check("23d 单波段 uint16 无 CRS 缩略图", gray_ds.raster is not None
          and (gray_ds.raster.preview_image or "").startswith("data:image/png"))
    bad_tif = tif_dir / "bad.tif"
    bad_tif.write_bytes(b"not a tiff")
    bad_ds = df._probe_tif(bad_tif, "bad.tif", 0)
    check("23e 损坏 TIFF 标 unavailable", bad_ds.status == DatasetStatus.unavailable and bool(bad_ds.reason))

    # ── 24. 重启恢复：workspace 物化 JSON 恢复为 Table ─────────
    DATA_REGISTRY._layers.clear()
    DATA_REGISTRY._geojson.clear()
    DATA_REGISTRY._schemas.clear()
    DATA_REGISTRY.load_workspace_layer_files(LAYER_DATA_DIR)
    lyr24 = DATA_REGISTRY.get_layer("tab17")
    check("24a workspace 恢复为 Table",
          lyr24 is not None and lyr24.geometry == GeometryType.table and lyr24.crs == "N/A" and lyr24.lazy)
    check("24b 恢复后 fields 非空可过滤", lyr24 is not None and len(lyr24.fields) > 0)

    # ── 25. 空间工具拒 Table（不得 stable_rate 演示数值）────────
    from app.tools.analysis import run_buffer, run_density, run_spatial_join  # noqa: E402
    from app.tools.coverage import run_coverage  # noqa: E402

    c = run_coverage({"layer_id": "tab17"})
    b = run_buffer({"layer_id": "tab17"})
    d = run_density({"layer_id": "tab17"})
    sj = run_spatial_join({"target_layer_id": "tab17", "overlay_layer_id": "tab17"})
    all_fail = not c.success and not b.success and not d.success and not sj.success
    no_demo = all("演示数值" not in (m.message or "") for m in (c, b, d, sj))
    check("25a 空间工具拒 Table（success=False）", all_fail,
          f"c={c.success} b={b.success} d={d.success} sj={sj.success}")
    check("25b 无演示数值兜底", no_demo)

    # ── 33. Table CSV 全量路径单次遍历（不先 list(reader) 再二次复制）────
    single_dir = Path(_FIX) / "single"
    single_dir.mkdir(parents=True, exist_ok=True)
    single_csv = single_dir / "t.csv"
    single_csv.write_text(
        "a,b\n" + "\n".join(f"r{i},{i * 10}" for i in range(500)), encoding="utf-8"
    )
    fc_full = df._table_csv_to_fc_full(single_csv)
    check("33a 表格 CSV 全量 FC 要素数与行数一致", len(fc_full.get("features", [])) == 500,
          f"n={len(fc_full.get('features', []))}")
    prop_ok = any(f.get("properties", {}).get("b") == "4990" for f in fc_full.get("features", []))
    check("33b 全量路径属性值原样保留（单次遍历未丢字段）", prop_ok)

    # ── 34. 行数超阈值但压缩文件小的 XLSX 标仅可查询 ────────────
    import openpyxl as _oxl

    old_table_thr = df._ENABLE_MAX_FEATURES_TABLE
    df._ENABLE_MAX_FEATURES_TABLE = 100  # 收紧表格行数阈值（运行时生效，扫描读模块全局）
    try:
        rowxl_dir = Path(_FIX) / "rowxl"
        rowxl_dir.mkdir(parents=True, exist_ok=True)
        wb = _oxl.Workbook()
        ws = wb.active
        ws.append(["name", "val"])
        for i in range(300):  # 300 > 100 行，但文件本身很小
            ws.append([f"p{i}", i])
        rowxl_path = rowxl_dir / "manyrows.xlsx"
        wb.save(rowxl_path)
        rowxl_src = df.link_source(str(rowxl_dir), alias="行数源")
        rowxl_ds = next((d for d in rowxl_src.datasets if d.dataset_id == "manyrows"), None)
        check("34a 小压缩包、超行数 XLSX 标仅可查询",
              rowxl_ds is not None and rowxl_ds.status == DatasetStatus.query_only,
              f"status={rowxl_ds.status if rowxl_ds else None} fc={rowxl_ds.feature_count if rowxl_ds else None}")
        small_wb = _oxl.Workbook()
        sws = small_wb.active
        sws.append(["name", "val"])
        sws.append(["a", 1])
        small_path = rowxl_dir / "small.xlsx"
        small_wb.save(small_path)
        rowxl_src2 = df.link_source(str(rowxl_dir), alias="行数源2")
        small_ds = next((d for d in rowxl_src2.datasets if d.dataset_id == "small"), None)
        check("34b 小行数 XLSX 仍可直接启用（不误伤）",
              small_ds is not None and small_ds.status == DatasetStatus.enableable,
              f"status={small_ds.status if small_ds else None}")
    finally:
        df._ENABLE_MAX_FEATURES_TABLE = old_table_thr

    # ── 35. 单个预计超过预算的候选不突破预算（首项也受限）────────
    old_budget = df._ENABLE_ALL_MAX_BYTES
    old_thr35 = df._ENABLE_MAX_FEATURES_TABLE
    df._ENABLE_MAX_FEATURES_TABLE = 100000  # 让 1000 行表格走 enableable（focus 预算拦截）
    df._ENABLE_ALL_MAX_BYTES = 50000  # 1000 行表格估算 1000*250=250000 > 50000 → 首项即被拦截
    try:
        one_dir = Path(_FIX) / "one"
        one_dir.mkdir(parents=True, exist_ok=True)
        (one_dir / "big.csv").write_text(
            "name,val\n" + "\n".join(f"p{i},{i}" for i in range(1000)), encoding="utf-8"
        )
        one_src = df.link_source(str(one_dir), alias="单项源")
        r_one = df.enable_all_datasets(one_src.source_id)
        check("35a 单个超预算候选不突破预算",
              r_one.get("planned_count") == 0 and r_one.get("enabled") == 0
              and r_one.get("capped_by_bytes") is True,
              f"planned={r_one.get('planned_count')} capped_by_bytes={r_one.get('capped_by_bytes')}")
    finally:
        df._ENABLE_ALL_MAX_BYTES = old_budget
        df._ENABLE_MAX_FEATURES_TABLE = old_thr35

    # ── 36. uploaded 与 folder 同名不覆盖，disable 只删 folder 图层 ──
    from app.models.layer import LayerSource  # noqa: E402
    from app.upload_service import register_feature_collection as _rfc  # noqa: E402

    _rfc(
        {"type": "FeatureCollection", "features": [
            {"type": "Feature", "geometry": {"type": "Point", "coordinates": [114.0, 22.5]}, "properties": {"n": 1}, "id": 0},
        ]},
        layer_id="dup",
        name="dup-uploaded",
        source=LayerSource.uploaded,
    )
    dup_dir = Path(_FIX) / "dup"
    dup_dir.mkdir(parents=True, exist_ok=True)
    (dup_dir / "dup.csv").write_text("name,lon,lat\np,114.0,22.5\n", encoding="utf-8")
    dup_src = df.link_source(str(dup_dir), alias="dup源")
    dup_ds = next((d for d in dup_src.datasets if d.dataset_id == "dup"), None)
    r_dup = df.enable_dataset(dup_src.source_id, "dup")
    folder_lid = r_dup.get("layer_id")
    check("36a folder 启用不与 uploaded 同名覆盖（得到稳定后缀 id）",
          folder_lid is not None and folder_lid != "dup" and folder_lid.startswith("dup_"),
          f"folder_lid={folder_lid}")
    up_layer = DATA_REGISTRY.get_layer("dup")
    folder_layer = DATA_REGISTRY.get_layer(folder_lid) if folder_lid else None
    check("36b uploaded 与 folder 两个图层同时存在",
          up_layer is not None and folder_layer is not None
          and up_layer.source == LayerSource.uploaded and folder_layer.source == LayerSource.folder)
    r_dup_d = df.disable_dataset(dup_src.source_id, "dup")
    check("36c 停用后 folder 图层移除", folder_lid is None or DATA_REGISTRY.get_layer(folder_lid) is None)
    check("36d uploaded 图层不受停用影响", DATA_REGISTRY.get_layer("dup") is not None)

    # ── 41. 冲突后缀 ID 的目录重扫后仍识别为已启用 ─────────────
    r_dup2 = df.enable_dataset(dup_src.source_id, "dup")  # uploaded 'dup' 仍在 → 仍走稳定后缀 id
    folder_lid2 = r_dup2.get("layer_id")
    df.rescan_source(dup_src.source_id)
    dup_src2 = df.get_source(dup_src.source_id)
    dup_ds2 = next((d for d in dup_src2.datasets if d.dataset_id == "dup"), None)
    check("41a 重扫后后缀图层仍识别为已启用",
          dup_ds2 is not None and dup_ds2.enabled is True and dup_ds2.layer_id == folder_lid2,
          f"enabled={dup_ds2.enabled if dup_ds2 else None} lid={dup_ds2.layer_id if dup_ds2 else None}")

    # ── 37. 侧车 meta.id 冲突拒绝 + 实际 geometry 以 GeoJSON 为准 ──
    from app.models.layer import GeometryType as _GT  # noqa: E402

    sidecar_dir = LAYER_DATA_DIR
    real_gj = {"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [114.0, 22.5]}, "properties": {"v": 1}},
    ]}
    (sidecar_dir / "sc_point.json").write_text(json.dumps(real_gj), encoding="utf-8")
    # 错位侧车：id 与文件 stem 不符 → 拒绝，走 geojson 兜底重建
    (sidecar_dir / "sc_point.meta.json").write_text(
        json.dumps({"id": "sc_WRONG", "name": "x", "source": "uploaded", "geometry": "Polygon",
                    "feature_count": 99, "fields": [], "data_path": "bad"}), encoding="utf-8")
    # 正确侧车但 geometry/feature_count 与实际 GeoJSON 冲突 → 派生字段以实际为准
    (sidecar_dir / "sc_good.json").write_text(json.dumps(real_gj), encoding="utf-8")
    (sidecar_dir / "sc_good.meta.json").write_text(
        json.dumps({"id": "sc_good", "name": "good", "source": "folder", "geometry": "Polygon",
                    "feature_count": 99, "fields": [], "data_path": "/tmp/x.csv", "rel_path": "x.csv",
                    "data_type": "csv", "visible": False, "lazy": True}), encoding="utf-8")
    DATA_REGISTRY.load_workspace_layer_files(sidecar_dir)
    sc_point = DATA_REGISTRY.get_layer("sc_point")
    check("37a meta.id 与 stem 不符 → 拒绝侧车、走 geojson 兜底",
          sc_point is not None and sc_point.source == LayerSource.uploaded
          and sc_point.geometry == _GT.point and sc_point.feature_count == 1,
          f"source={sc_point.source if sc_point else None} geom={sc_point.geometry if sc_point else None}")
    sc_good = DATA_REGISTRY.get_layer("sc_good")
    check("37b 侧车正确：注册语义从侧车恢复（source/data_type/lazy/visible/data_path/rel_path）",
          sc_good is not None and sc_good.source == LayerSource.folder
          and sc_good.data_type == "csv" and sc_good.lazy is True and sc_good.visible is False
          and sc_good.data_path == "/tmp/x.csv" and sc_good.rel_path == "x.csv",
          f"src={sc_good.source if sc_good else None} type={sc_good.data_type if sc_good else None}")
    check("37c 数据派生字段以实际 GeoJSON 为准（geometry=Point、fc=1，非侧车 Polygon/99）",
          sc_good is not None and sc_good.geometry == _GT.point and sc_good.feature_count == 1,
          f"geom={sc_good.geometry if sc_good else None} fc={sc_good.feature_count if sc_good else None}")

    # ── 38. 项目详情不列出 .meta.json ──────────────────────────
    from app.main import api_project_detail as _api_detail  # noqa: E402
    from app.projects import create_project as _create_project  # noqa: E402
    from app.projects import project_layer_dir as _pld  # noqa: E402

    proj = _create_project("meta测试项目")
    pid = proj["id"]
    pdir = _pld(pid)
    (pdir / "a.json").write_text(json.dumps(real_gj), encoding="utf-8")
    (pdir / "a.meta.json").write_text(json.dumps({"id": "a"}), encoding="utf-8")
    detail = _api_detail(pid)
    names = [f["name"] for f in detail.files]
    check("38a 项目详情不列出 .meta.json",
          "a.json" in names and "a.meta.json" not in names,
          f"names={names}")

    # ── 39. 抽稀样本 query_features 不返回虚假全量总数 ─────────
    from app.tools.query_features import run_query_features as _rqf  # noqa: E402

    q_s = _rqf({"layer_id": "bigtab", "filter": {"name": "p1"}})
    check("39a sampled 图层 query 返回 sample_matched/sample_size 而非 total_matched",
          q_s.success and q_s.output.get("sampled") is True
          and "total_matched" not in q_s.output
          and "sample_matched" in q_s.output and "sample_size" in q_s.output,
          f"keys={sorted(q_s.output.keys()) if q_s.success else q_s.message}")

    # ── 40. data_type/source/data_path/rel_path/visible 注册、会话、重启恢复一致 ──
    lyr40 = DATA_REGISTRY.get_layer("tab17")
    check("40a 重启恢复后 source=folder",
          lyr40 is not None and lyr40.source == LayerSource.folder,
          f"source={lyr40.source if lyr40 else None}")
    check("40b 恢复后 data_path/rel_path/data_type/visible 保留",
          lyr40 is not None and lyr40.data_path and lyr40.rel_path == "tab17.csv"
          and lyr40.data_type == "csv" and lyr40.visible is True,
          f"dp={lyr40.data_path if lyr40 else None} rp={lyr40.rel_path if lyr40 else None} "
          f"dt={lyr40.data_type if lyr40 else None} visible={lyr40.visible if lyr40 else None}")

    # ── 42. Table 批量启用内存估算保守下限 ─────────────────────
    from app.models.data_folder import FolderDataset as _FD  # noqa: E402

    d_wide = _FD(dataset_id="w", name="w", rel_path="w.csv", format="csv", size=400_000,
                 geometry="table", feature_count=2, status=DatasetStatus.enableable)
    est_wide = df._estimate_load_bytes(d_wide)
    check("42a 少行超宽 enableable Table 不被严重低估（≥文件字节）",
          est_wide >= 400_000, f"est={est_wide}")

    d_qo = _FD(dataset_id="q", name="q", rel_path="q.csv", format="csv", size=5_000_000,
               geometry="table", feature_count=1_000_000, status=DatasetStatus.query_only)
    est_qo = df._estimate_load_bytes(d_qo)
    # query_only 抽样 min(1e6, 5e4)=5e4 行 × per_row(250) = 12.5MB，且 ≥ 文件字节 5MB
    check("42b query_only 宽表用 Table 专属估算（抽样行数×每行，且 ≥ 文件字节）",
          est_qo >= 5_000_000 and est_qo == max(50_000 * 250, 5_000_000), f"est={est_qo}")

    # c. query_only Table 单项超预算仍 planned_count=0（集成）
    old_budget42 = df._ENABLE_ALL_MAX_BYTES
    old_thr42 = df._ENABLE_MAX_FEATURES_TABLE
    df._ENABLE_MAX_FEATURES_TABLE = 100  # 300 行无坐标 CSV → query_only
    df._ENABLE_ALL_MAX_BYTES = 1000  # 极小预算：300 行抽样估算 ≥ 300*250=75000 远超
    try:
        qo_dir = Path(_FIX) / "qo"
        qo_dir.mkdir(parents=True, exist_ok=True)
        (qo_dir / "qotab.csv").write_text(
            "name,val\n" + "\n".join(f"p{i},{i}" for i in range(300)), encoding="utf-8"
        )
        qo_src = df.link_source(str(qo_dir), alias="qo源")
        r_qo = df.enable_all_datasets(qo_src.source_id)
        check("42c query_only Table 单项超预算仍 planned_count=0",
              r_qo.get("planned_count") == 0 and r_qo.get("capped_by_bytes") is True,
              f"planned={r_qo.get('planned_count')} capped_by_bytes={r_qo.get('capped_by_bytes')}")
    finally:
        df._ENABLE_ALL_MAX_BYTES = old_budget42
        df._ENABLE_MAX_FEATURES_TABLE = old_thr42

    # ── 43. bbox 随真实 GeoJSON 更新 ──────────────────────────
    # a. 侧车 bbox 错误、GeoJSON 无顶层 bbox → 恢复为实际 bbox
    (sidecar_dir / "sc_bbox.json").write_text(
        json.dumps({"type": "FeatureCollection", "features": [
            {"type": "Feature", "geometry": {"type": "Point", "coordinates": [114.0, 22.5]}, "properties": {}},
        ]}), encoding="utf-8")
    (sidecar_dir / "sc_bbox.meta.json").write_text(
        json.dumps({"id": "sc_bbox", "source": "uploaded", "geometry": "Point",
                    "feature_count": 1, "fields": [], "bbox": [0.0, 0.0, 0.0, 0.0]}), encoding="utf-8")
    DATA_REGISTRY.load_workspace_layer_files(sidecar_dir)
    sc_bbox = DATA_REGISTRY.get_layer("sc_bbox")
    check("43a 侧车 bbox 错误、GeoJSON 无顶层 bbox → 恢复为实际 bbox",
          sc_bbox is not None and sc_bbox.bbox == [114.0, 22.5, 114.0, 22.5],
          f"bbox={sc_bbox.bbox if sc_bbox else None}")

    # b. 编辑几何后 bbox 更新
    from app.main import update_layer_geojson as _ugj  # noqa: E402
    from app.models.layer import UpdateLayerGeojsonRequest as _UGR  # noqa: E402

    new_gj = {"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [115.0, 23.5]}, "properties": {}},
    ]}
    upd = _ugj("sc_point", _UGR(geojson=new_gj))
    check("43b 编辑几何后 bbox 更新",
          upd.bbox == [115.0, 23.5, 115.0, 23.5], f"bbox={upd.bbox}")

    # c. Table 恢复后 bbox 为 None
    tab_bbox = DATA_REGISTRY.get_layer("tab17")
    check("43c Table 恢复后 bbox 为 None", tab_bbox is not None and tab_bbox.bbox is None,
          f"bbox={tab_bbox.bbox if tab_bbox else None}")

    # ── 44. 无坐标 Table CSV 逻辑行数统计 ─────────────────────
    log_dir = Path(_FIX) / "logical"
    log_dir.mkdir(parents=True, exist_ok=True)
    (log_dir / "two.csv").write_text("a,b\n1,2\n3,4\n", encoding="utf-8")
    (log_dir / "multi.csv").write_text('a,b\n"line1\nline2",x\n', encoding="utf-8")
    log_ds = {d.dataset_id: d for d in df.scan_folder(log_dir)}
    check("44a 表头 + 2 行数据 feature_count=2",
          log_ds.get("two") is not None and log_ds["two"].feature_count == 2,
          f"fc={log_ds.get('two').feature_count if log_ds.get('two') else None}")
    check("44b 引号内换行按一条逻辑记录计数",
          log_ds.get("multi") is not None and log_ds["multi"].feature_count == 1,
          f"fc={log_ds.get('multi').feature_count if log_ds.get('multi') else None}")

    thr_dir = Path(_FIX) / "thr"
    thr_dir.mkdir(parents=True, exist_ok=True)
    (thr_dir / "exact.csv").write_text("a,b\n1,2\n3,4\n", encoding="utf-8")  # 2 行数据
    (thr_dir / "over.csv").write_text("a,b\n1,2\n3,4\n5,6\n", encoding="utf-8")  # 3 行数据
    old_thr44 = df._ENABLE_MAX_FEATURES_TABLE
    df._ENABLE_MAX_FEATURES_TABLE = 2
    try:
        thr_ds = {d.dataset_id: d for d in df.scan_folder(thr_dir)}
        check("44c 阈值边界无 off-by-one（2 行 enableable、3 行 query_only）",
              thr_ds.get("exact") is not None and thr_ds["exact"].status == DatasetStatus.enableable
              and thr_ds.get("over") is not None and thr_ds["over"].status == DatasetStatus.query_only
              and thr_ds["exact"].feature_count == 2 and thr_ds["over"].feature_count == 3,
              f"exact={thr_ds.get('exact').status if thr_ds.get('exact') else None} "
              f"over={thr_ds.get('over').status if thr_ds.get('over') else None}")
    finally:
        df._ENABLE_MAX_FEATURES_TABLE = old_thr44

    # ── 45. Agent folder_disable_dataset 用实际 removed_layer_id ──
    from app.models.events import EventType  # noqa: E402
    from app.tools import folder_ops as _fo  # noqa: E402
    from app.tools.context import ToolContext, reset_current_context, set_current_context  # noqa: E402

    class _Rec:
        def __init__(self):
            self.events = []
        def emit(self, et, data):
            self.events.append((et, data))

    rec = _Rec()
    tok = set_current_context(ToolContext(event_emitter=rec))
    try:
        r_tool = _fo.run_folder_disable_dataset({"source_id": dup_src.source_id, "dataset_id": "dup"})
    finally:
        reset_current_context(tok)
    check("45a 工具停用返回实际 removed_layer_id（后缀 id）",
          r_tool.success and r_tool.output.get("removed_layer_id") == folder_lid2,
          f"out={r_tool.output}")
    check("45b 事件 payload 用实际 removed_layer_id",
          any(e[0] == EventType.LAYER_REMOVED and e[1].get("layer_id") == folder_lid2 for e in rec.events),
          f"events={rec.events}")
    check("45c uploaded「dup」图层不受 Agent 停用影响", DATA_REGISTRY.get_layer("dup") is not None)

    rec2 = _Rec()
    tok2 = set_current_context(ToolContext(event_emitter=rec2))
    try:
        r_tool2 = _fo.run_folder_disable_dataset({"source_id": dup_src.source_id, "dataset_id": "dup"})
    finally:
        reset_current_context(tok2)
    check("45d 幂等停用不发送伪删除事件",
          r_tool2.success and r_tool2.output.get("already_disabled") is True
          and not any(e[0] == EventType.LAYER_REMOVED for e in rec2.events),
          f"out={r_tool2.output} events={rec2.events}")

    # ── 46. folder_list 指定不存在 source 返回明确失败 ─────────
    r_missing = _fo.run_folder_list({"source_id": "nonexistent_xyz"})
    check("46a folder_list 指定不存在 source 返回失败",
          r_missing.success is False and "不存在" in (r_missing.message or ""),
          f"msg={r_missing.message}")

    print(f"\n结果：{TOTAL - len(FAILS)}/{TOTAL} 通过" + (f"，失败：{FAILS}" if FAILS else ""))
    return 1 if FAILS else 0


if __name__ == "__main__":
    code = main()
    shutil.rmtree(_WS, ignore_errors=True)
    shutil.rmtree(_FIX, ignore_errors=True)
    sys.exit(code)
